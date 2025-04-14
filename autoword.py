def gerar_documento():
    from docx import Document

    documento = Document()
    documento.add_heading('Documento Gerado com Python', level=0)

    paragrafo = documento.add_paragraph('\nEsse é um parágrafo.')
    paragrafo.add_run('\nEsse texto está em negrito.').bold = True
    paragrafo.add_run('\nEsse texto está em itálico.').italic = True

    qtd_pessoas = int(input("Quantas pessoas deseja adicionar à tabela? "))
    table = documento.add_table(rows=qtd_pessoas + 1, cols=2)
    table.cell(0, 0).text = 'Nome'
    table.cell(0, 1).text = 'Idade'

    for i in range(1, qtd_pessoas + 1):
        nome = input(f"Digite o nome da pessoa {i}: ")
        idade = input(f"Digite a idade de {nome}: ")
        table.cell(i, 0).text = nome
        table.cell(i, 1).text = idade

    nome_arquivo = input("Digite o nome do arquivo (sem extensão): ")
    documento.save(f"{nome_arquivo}.docx")
    print(f"\nDocumento '{nome_arquivo}.docx' salvo com sucesso!")


# Execução principal
print("Bem-vindo ao AutoWord!")
resposta = input("Deseja gerar um documento de exemplo? (1 - SIM / 2 - NÃO): ")

if resposta == "1":
    gerar_documento()
else:
    print("Programa encerrado.")
